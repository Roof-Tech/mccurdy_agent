"""Turn a prospect list into ready-to-send follow-up email drafts (one sequence per prospect type).

Drafts are written to a CSV (for mail merge / your CRM) and one .md per prospect for review.
Nothing is sent automatically: you (or Claude with Gmail, after you approve) send them.
"""
import csv
import datetime as dt
import re
from pathlib import Path

from openpyxl import load_workbook

from .config import ROOT, load_config

TEMPLATES = ROOT / "templates" / "emails"
STEP_RE = re.compile(r"^## Step (\d+) — Day (\d+) — Subject: (.+)$", re.M)


def load_sequence(segment: str) -> list[dict]:
    text = (TEMPLATES / f"{segment}.md").read_text(encoding="utf-8")
    parts = STEP_RE.split(text)
    steps = []
    for i in range(1, len(parts), 4):
        steps.append({"step": int(parts[i]), "day": int(parts[i + 1]), "subject": parts[i + 2].strip(),
                      "body": parts[i + 3].strip()})
    return steps


def read_prospects(path: str | Path) -> list[dict]:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))
    ws = load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr_idx = next(i for i, r in enumerate(rows)
                   if r and ("Score" in r or "email" in [str(x).lower() for x in r if x]))
    headers = [str(h or "") for h in rows[hdr_idx]]
    labels = {"Owner": "owner", "Property address": "site_address", "City": "city", "Roof sq ft (est.)": "building_sqft",
              "Business / Building": "name", "Score": "score", "Use": "use"}
    keys = [labels.get(h, h.lower().replace(" ", "_")) for h in headers]
    return [dict(zip(keys, r)) for r in rows[hdr_idx + 1:] if any(r)]


def _tidy(v) -> str:
    """Assessor data is often ALL CAPS; emails read better in normal case."""
    v = str(v or "").strip()
    return re.sub(r"\b(Ne|Nw|Se|Sw)\b", lambda m: m.group(0).upper(), v.title()) if v.isupper() else v


class _Safe(dict):
    def __missing__(self, key):
        return "{" + key + "}"


def personalize(text: str, p: dict, cfg: dict) -> str:
    co = cfg["company"]
    sqft = float(re.sub(r"[^0-9.]", "", str(p.get("building_sqft") or 0)) or 0)
    contact = str(p.get("contact_name") or p.get("first_name") or "").strip()
    first = contact.split()[0].title() if contact else "there"
    signature = (f"{co['signer_name']}\n{co['signer_title']}, {co['name']}\n{co['phone']} · {co['email']}\n"
                 f"CSLB #{co['license']}")
    values = _Safe(
        first_name=first,
        site_address=_tidy(p.get("site_address") or p.get("name")) or "your building",
        city=_tidy(p.get("city")) or "the Central Valley",
        building_sqft=f"{sqft:,.0f}" if sqft else "your",
        sender_name=co["signer_name"],
        company_name=co["name"],
        signature=signature,
    )
    return text.format_map(values)


def build_drafts(prospects: list[dict], segment: str, out_dir: str | Path, start: dt.date | None = None,
                 cfg: dict | None = None) -> Path:
    cfg = cfg or load_config()
    start = start or dt.date.today()
    steps = load_sequence(segment)
    out_dir = Path(out_dir)
    (out_dir / "drafts").mkdir(parents=True, exist_ok=True)
    csv_path = out_dir / f"followups_{segment}.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["send_on", "to_email", "contact", "property", "step", "subject", "body"])
        for n, p in enumerate(prospects, 1):
            md = []
            for s in steps:
                subject = personalize(s["subject"], p, cfg)
                body = personalize(s["body"], p, cfg)
                send_on = start + dt.timedelta(days=s["day"])
                w.writerow([send_on.isoformat(), p.get("email", ""), p.get("contact_name", ""),
                            p.get("site_address", ""), s["step"], subject, body])
                md.append(f"## {send_on:%a %b %d} — Step {s['step']}\n**Subject:** {subject}\n\n{body}\n")
            slug = re.sub(r"[^a-z0-9]+", "-", str(p.get("site_address") or p.get("name") or n).lower()).strip("-")
            (out_dir / "drafts" / f"{n:03d}-{slug[:50]}.md").write_text("\n".join(md), encoding="utf-8")
    return csv_path
