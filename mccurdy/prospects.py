"""Find and rank commercial flat-roof prospects.

Two sources, both free:
  1. County assessor exports (CSV) — owner names and mailing addresses. Download the parcel/roll file
     from your county assessor (or a data vendor) and run `prospects score <file.csv>`.
  2. OpenStreetMap building footprints — large warehouse/industrial/commercial roofs by city or county,
     with roof size estimated from the footprint. Run `prospects map --area "Fresno"`.
Both write a branded Excel sheet sorted best-first.
"""
import csv
import datetime as dt
import json
import math
import re
import urllib.parse
import urllib.request
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import load_config

# Header names seen in California assessor exports, mapped to our field names.
COLUMN_ALIASES = {
    "apn": ["apn", "parcel", "parcel_number", "parcel_no", "assessor_parcel_number", "ain"],
    "owner": ["owner", "owner_name", "owner1", "ownername", "assessee", "assessee_name"],
    "site_address": ["situs", "situs_address", "site_address", "property_address", "address", "situs_addr"],
    "city": ["situs_city", "city", "site_city"],
    "mail_address": ["mail_address", "mailing_address", "mailing_addr", "owner_address", "mail_addr"],
    "use": ["use_code_description", "land_use_description", "use_description", "use_desc", "land_use", "landuse",
            "property_use", "property_type", "use_code", "use"],
    "building_sqft": ["building_sqft", "bldg_sqft", "building_area", "improvement_sqft", "sqft", "living_area",
                      "bldg_area", "gross_building_area"],
    "year_built": ["year_built", "yr_built", "yearbuilt", "effective_year"],
}

PRIORITY_USES = ["warehouse", "distribution", "industrial", "manufactur", "storage", "cold storage", "packing",
                 "processing", "logistics"]
COMMERCIAL_USES = PRIORITY_USES + ["commercial", "retail", "office", "shopping", "store", "market", "church",
                                   "school", "medical", "hotel", "motel", "restaurant", "auto", "service",
                                   "multi", "apartment", "mixed"]
INVESTOR_MARKERS = re.compile(r"\b(LLC|L\.L\.C|INC|CORP|LP|L\.P|TRUST|PROPERTIES|HOLDINGS|INVESTMENTS|REIT|PARTNERS|"
                              r"REALTY|CAPITAL|VENTURES|GROUP|FUND|ASSOCIATES)\b", re.I)


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", h.strip().lower()).strip("_")


def map_columns(headers: list[str]) -> dict:
    normed = {_norm(h): h for h in headers}
    found = {}
    for field, aliases in COLUMN_ALIASES.items():
        for a in aliases:
            if a in normed:
                found[field] = normed[a]
                break
    return found


def _num(v) -> float:
    try:
        return float(re.sub(r"[^0-9.]", "", str(v)) or 0)
    except ValueError:
        return 0.0


def score_parcel(p: dict, min_sqft: float, this_year: int | None = None) -> tuple[int, list[str]]:
    """0–100 lead score with plain-English reasons."""
    this_year = this_year or dt.date.today().year
    score, why = 0, []
    sqft = p.get("building_sqft", 0)
    use = str(p.get("use", "")).lower()

    if sqft >= 100_000:
        score += 40; why.append("100k+ sq ft roof")
    elif sqft >= 50_000:
        score += 32; why.append("50k+ sq ft roof")
    elif sqft >= 20_000:
        score += 22; why.append("20k+ sq ft roof")
    elif sqft >= min_sqft:
        score += 12

    if any(u in use for u in PRIORITY_USES):
        score += 25; why.append("warehouse/industrial (almost always low-slope)")
    elif any(u in use for u in COMMERCIAL_USES):
        score += 15; why.append("commercial use")

    yb = int(p.get("year_built") or 0)
    if yb:
        age = this_year - yb
        if age >= 15:
            score += 20; why.append(f"built {yb} — original or aging roof likely")
        elif age >= 8:
            score += 10; why.append(f"built {yb}")

    if INVESTOR_MARKERS.search(str(p.get("owner", ""))):
        score += 10; why.append("owned by a business/investor entity")
    site, mail = str(p.get("site_address", "")).lower(), str(p.get("mail_address", "")).lower()
    if mail and site and site.split(" ")[0:2] != mail.split(" ")[0:2]:
        score += 5; why.append("absentee owner (mails elsewhere)")
    return min(score, 100), why


def score_assessor_csv(path: str | Path, min_sqft: float = 10_000, uses: str = "commercial") -> list[dict]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        cols = map_columns(reader.fieldnames or [])
        if "building_sqft" not in cols:
            raise ValueError(f"Could not find a building square-feet column in {path}. Columns: {reader.fieldnames}")
        results = []
        for row in reader:
            p = {field: row.get(src, "") for field, src in cols.items()}
            p["building_sqft"] = _num(p.get("building_sqft"))
            p["year_built"] = int(_num(p.get("year_built"))) or ""
            use = str(p.get("use", "")).lower()
            if p["building_sqft"] < min_sqft:
                continue
            if uses == "commercial" and "use" in cols and not any(u in use for u in COMMERCIAL_USES):
                continue
            p["score"], reasons = score_parcel(p, min_sqft)
            p["why"] = "; ".join(reasons)
            p["source"] = Path(path).name
            results.append(p)
    results.sort(key=lambda r: (-r["score"], -r["building_sqft"]))
    return results


# ── OpenStreetMap ────────────────────────────────────────────

OVERPASS_URL = "https://overpass-api.de/api/interpreter"
SQFT_PER_SQM = 10.7639


def overpass_query(area: str, min_sqft: float) -> str:
    kinds = "warehouse|industrial|commercial|retail|manufacture|supermarket|office|storage_tank"
    return f"""[out:json][timeout:120];
area["name"="{area}"]["boundary"="administrative"]->.a;
way["building"~"^({kinds})$"](area.a);
out tags geom;"""


def polygon_area_sqm(coords: list[tuple[float, float]]) -> float:
    """Shoelace area on a local equirectangular projection — accurate to a few % for building footprints."""
    if len(coords) < 3:
        return 0.0
    lat0 = math.radians(sum(c[0] for c in coords) / len(coords))
    r = 6_371_008.8
    pts = [(math.radians(lon) * r * math.cos(lat0), math.radians(lat) * r) for lat, lon in coords]
    a = 0.0
    for (x1, y1), (x2, y2) in zip(pts, pts[1:] + pts[:1]):
        a += x1 * y2 - x2 * y1
    return abs(a) / 2


def parse_overpass(data: dict, min_sqft: float) -> list[dict]:
    out = []
    for el in data.get("elements", []):
        geom = el.get("geometry") or []
        coords = [(g["lat"], g["lon"]) for g in geom]
        sqft = polygon_area_sqm(coords) * SQFT_PER_SQM
        if sqft < min_sqft:
            continue
        tags = el.get("tags", {})
        lat = sum(c[0] for c in coords) / len(coords)
        lon = sum(c[1] for c in coords) / len(coords)
        addr = " ".join(x for x in [tags.get("addr:housenumber", ""), tags.get("addr:street", "")] if x)
        kind = tags.get("building", "")
        p = {
            "name": tags.get("name", ""),
            "site_address": addr,
            "city": tags.get("addr:city", ""),
            "use": kind,
            "building_sqft": round(sqft),
            "lat": round(lat, 6),
            "lon": round(lon, 6),
            "map_link": f"https://www.google.com/maps/@{lat:.6f},{lon:.6f},120m/data=!3m1!1e3",
            "osm_id": el.get("id"),
            "source": "OpenStreetMap",
        }
        p["score"], reasons = score_parcel(p, min_sqft)
        p["why"] = "; ".join(reasons)
        out.append(p)
    out.sort(key=lambda r: (-r["score"], -r["building_sqft"]))
    return out


def fetch_osm_buildings(area: str, min_sqft: float = 20_000) -> list[dict]:
    body = urllib.parse.urlencode({"data": overpass_query(area, min_sqft)}).encode()
    req = urllib.request.Request(OVERPASS_URL, data=body, headers={"User-Agent": "mccurdy-prospects/1.0"})
    with urllib.request.urlopen(req, timeout=180) as resp:
        return parse_overpass(json.load(resp), min_sqft)


# ── Output ───────────────────────────────────────────────────

COLUMNS = [("score", "Score", 8), ("why", "Why it's a fit", 48), ("name", "Business / Building", 26),
           ("owner", "Owner", 30), ("site_address", "Property address", 32), ("city", "City", 14),
           ("use", "Use", 18), ("building_sqft", "Roof sq ft (est.)", 15), ("year_built", "Year built", 10),
           ("mail_address", "Owner mailing address", 32), ("apn", "APN", 16), ("map_link", "Aerial view", 18),
           ("status", "Status", 14), ("next_step", "Next step", 22), ("notes", "Notes", 30), ("source", "Source", 16)]


def write_prospects_xlsx(rows: list[dict], out_path: str | Path, title: str, cfg: dict | None = None) -> Path:
    cfg = cfg or load_config()
    brand = cfg["brand"]
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    cols = [c for c in COLUMNS if c[0] in ("status", "next_step", "notes") or any(r.get(c[0]) for r in rows)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Prospects"
    ws["A1"] = f"{cfg['company']['name']} — {title}"
    ws["A1"].font = Font(bold=True, size=14, color=brand["blue"].lstrip("#"))
    ws["A2"] = f"{len(rows)} properties · generated {dt.date.today():%b %d, %Y} · sorted best-first"
    ws["A2"].font = Font(italic=True, color=brand["gray"].lstrip("#"))
    header_fill = PatternFill("solid", fgColor=brand["blue"].lstrip("#"))
    hot_fill = PatternFill("solid", fgColor=brand["green_light"].lstrip("#"))
    for j, (_, label, width) in enumerate(cols, 1):
        cell = ws.cell(row=4, column=j, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = width
    for i, r in enumerate(rows, 5):
        for j, (key, _, _) in enumerate(cols, 1):
            v = r.get(key, "")
            cell = ws.cell(row=i, column=j, value=v)
            if key == "map_link" and v:
                cell.hyperlink = v
                cell.value = "Open aerial"
                cell.font = Font(color=brand["blue_bright"].lstrip("#"), underline="single")
            if key == "building_sqft":
                cell.number_format = "#,##0"
            if r.get("score", 0) >= 70:
                cell.fill = hot_fill
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(4, len(rows) + 4)}"
    wb.save(out_path)
    return out_path
